import time
import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from sklearn.model_selection import cross_validate
from sklearn.datasets import load_iris
from sklearn.model_selection import train_test_split
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import LogisticRegression
from sklearn.tree import DecisionTreeClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.compose import ColumnTransformer, make_column_selector
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import OneHotEncoder
from sklearn.pipeline import Pipeline
from sklearn.metrics import (
    accuracy_score,
    precision_score,
    recall_score,
    f1_score,
    confusion_matrix,
    ConfusionMatrixDisplay,
)

def create_excel_file(dataframes):
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, df in dataframes.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]

            # Format the header
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid", fgColor="244062"
                )
                cell.alignment = Alignment(
                    horizontal="center"
                )

            # Adjust column widths
            for column in ws.columns:
                max_length = max(
                    len(str(cell.value or ""))
                    for cell in column
                )
                ws.column_dimensions[
                    get_column_letter(column[0].column)
                ].width = min(max_length + 3, 30)

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

    output.seek(0)
    return output.getvalue()



st.set_page_config(
    page_title="AI Model Comparison Tool",
    page_icon="🤖",
    layout="wide",
)

st.title("🤖 AI Model Comparison Tool")
st.write(
    "Train, evaluate, and compare machine learning models "
    "using a consistent dataset and evaluation process."
)

# Sidebar settings
st.sidebar.header("⚙️ Model Settings")

test_size = st.sidebar.slider(
    "Test data percentage",
    min_value=10,
    max_value=40,
    value=20,
    step=5,
)

tree_depth = st.sidebar.slider(
    "Decision Tree maximum depth",
    min_value=1,
    max_value=10,
    value=3,
)

forest_trees = st.sidebar.slider(
    "Random Forest number of trees",
    min_value=10,
    max_value=200,
    value=100,
    step=10,
)

# Load dataset
from data_utils import load_csv

# Dataset selection
st.sidebar.header("📂 Dataset")

dataset_option = st.sidebar.radio(
    "Choose your dataset:",
    ["Iris (Built-in)", "Upload CSV"],
)

if dataset_option == "Iris (Built-in)":
    iris = load_iris()
    df = pd.DataFrame(
        iris.data,
        columns=iris.feature_names,
    )
    df["target"] = iris.target
    target_column = "target"
    class_names = iris.target_names.tolist()

else:
    uploaded_file = st.sidebar.file_uploader(
        "Upload a CSV file",
        type=["csv"],
    )

    if uploaded_file is None:
        st.info("Please upload a CSV file to continue.")
        st.stop()

    df, error = load_csv(uploaded_file)

    if error:
        st.error(error)
        st.stop()

    st.subheader("👀 Dataset Preview")
    st.dataframe(df.head(), use_container_width=True)

    target_column = st.selectbox(
        "Select the target column to predict:",
        df.columns.tolist(),
    )

    class_names = sorted(
        df[target_column].dropna().astype(str).unique().tolist()
    )

X = df.drop(columns=[target_column])
y = df[target_column]

if y.isna().any():
    st.error("The target column contains missing values. Please clean it first.")
    st.stop()

if X.shape[1] == 0:
    st.error("The dataset must contain at least one feature column.")
    st.stop()
# Identify numeric and categorical columns
numeric_features = X.select_dtypes(
    include=["number"]
).columns.tolist()

categorical_features = X.select_dtypes(
    exclude=["number"]
).columns.tolist()

# Preprocess numeric and categorical data
numeric_pipeline = Pipeline([
    ("imputer", SimpleImputer(strategy="median")),
    ("scaler", StandardScaler()),
])

categorical_pipeline = Pipeline([
    ("imputer", SimpleImputer(strategy="most_frequent")),
    ("encoder", OneHotEncoder(handle_unknown="ignore")),
])

preprocessor = ColumnTransformer([
    ("numeric", numeric_pipeline, numeric_features),
    ("categorical", categorical_pipeline, categorical_features),
])

# Split the data
try:
    X_train, X_test, y_train, y_test = train_test_split(
        X,
        y,
        test_size=test_size / 100,
        random_state=42,
        stratify=y,
    )
except ValueError:
    X_train, X_test, y_train, y_test = train_test_split(
        X,
        y,
        test_size=test_size / 100,
        random_state=42,
    )

    # Validate target before training classification models
if y.nunique() < 2:
    st.error(
        "The target column must contain at least two different classes."
    )
    st.stop()

if pd.api.types.is_numeric_dtype(y):
    if y.nunique() > 20 and y.nunique() / len(y) > 0.05:
        st.error(
    "The selected target column contains continuous numerical values. "
    "The current models support classification only. "
    "Please select a categorical target column, such as Result (Pass/Fail), "
    "or use a classification dataset."
)
        st.stop()

# Define machine learning models
models = {
    "Logistic Regression": Pipeline([
        ("preprocessor", preprocessor),
        ("classifier", LogisticRegression(max_iter=1000)),
    ]),
    "Decision Tree": Pipeline([
        ("preprocessor", preprocessor),
        ("classifier", DecisionTreeClassifier(
            max_depth=tree_depth,
            random_state=42,
        )),
    ]),
    "Random Forest": Pipeline([
        ("preprocessor", preprocessor),
        ("classifier", RandomForestClassifier(
            n_estimators=forest_trees,
            random_state=42,
        )),
    ]),
}
results = []
trained_models = {}

# Train and evaluate
for name, model in models.items():
    start_train = time.perf_counter()
    model.fit(X_train, y_train)
    train_ms = (time.perf_counter() - start_train) * 1000

    start_predict = time.perf_counter()
    predictions = model.predict(X_test)
    predict_ms = (time.perf_counter() - start_predict) * 1000

    trained_models[name] = model

    results.append({
        "Model": name,
        "Accuracy": accuracy_score(y_test, predictions),
        "Precision": precision_score(
            y_test, predictions,
            average="weighted",
            zero_division=0,
        ),
        "Recall": recall_score(
            y_test, predictions,
            average="weighted",
            zero_division=0,
        ),
        "F1 Score": f1_score(
            y_test, predictions,
            average="weighted",
            zero_division=0,
        ),
        "Training Time (ms)": train_ms,
        "Prediction Time (ms)": predict_ms,
    })

results_df = pd.DataFrame(results)



# Cross-Validation
st.subheader("🔄 Cross-Validation Results")

cv_results = []

for name, model in models.items():
    scores = cross_validate(
        model,
        X_train,
        y_train,
        cv=5,
        scoring={
            "accuracy": "accuracy",
            "precision": "precision_weighted",
            "recall": "recall_weighted",
            "f1": "f1_weighted",
        },
        error_score="raise",
    )

    cv_results.append({
        "Model": name,
        "CV Accuracy": scores["test_accuracy"].mean(),
        "CV Precision": scores["test_precision"].mean(),
        "CV Recall": scores["test_recall"].mean(),
        "CV F1 Score": scores["test_f1"].mean(),
    })

cv_results_df = pd.DataFrame(cv_results)


st.dataframe(
    cv_results_df.style.format({
        "CV Accuracy": "{:.2%}",
        "CV Precision": "{:.2%}",
        "CV Recall": "{:.2%}",
        "CV F1 Score": "{:.2%}",
    }),
    use_container_width=True,
)

st.download_button(
    label="📥 Download Excel Results",
    data=create_excel_file({
        "Model Comparison": results_df
    }),
    file_name="model_comparison_results.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)


# Best Model Based on Cross-Validation
st.subheader("🏆 Best Performing Model")

best_model = cv_results_df.loc[
    cv_results_df["CV Accuracy"].idxmax()
]

st.success(
    f"Best Model: {best_model['Model']} "
    f"| Mean CV Accuracy: {best_model['CV Accuracy']:.2%}"
)

# Main results
st.subheader("📊 Model Performance")

metric = st.selectbox(
    "Rank models by:",
    ["Accuracy", "Precision", "Recall", "F1 Score"],
)

best_row = results_df.loc[results_df[metric].idxmax()]

col1, col2, col3 = st.columns(3)

col1.metric("Best Model", best_row["Model"])
col2.metric("Best " + metric, f"{best_row[metric]:.2%}")
col3.metric("Test Samples", len(y_test))

st.subheader("📋 Detailed Comparison")

st.dataframe(
    results_df.style.format({
        "Accuracy": "{:.2%}",
        "Precision": "{:.2%}",
        "Recall": "{:.2%}",
        "F1 Score": "{:.2%}",
        "Training Time (ms)": "{:.2f}",
        "Prediction Time (ms)": "{:.2f}",
    }),
    use_container_width=True,
    hide_index=True,
)
st.download_button(
    label="📥 Download Cross-Validation Excel",
    data=create_excel_file({
        "Cross-Validation": cv_results_df
    }),
    file_name="cross_validation_results.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)

# Performance chart
st.subheader("📈 Metric Comparison")

chart_data = results_df.set_index("Model")[
    ["Accuracy", "Precision", "Recall", "F1 Score"]
]

st.bar_chart(chart_data)

# Time comparison
st.subheader("⏱️ Execution Time Comparison")

time_data = results_df.set_index("Model")[
    ["Training Time (ms)", "Prediction Time (ms)"]
]

st.bar_chart(time_data)

# Confusion matrix
st.subheader("🎯 Confusion Matrix")

selected_name = st.selectbox(
    "Choose a model to inspect:",
    list(trained_models.keys()),
    key="confusion_model",
)

selected_model = trained_models[selected_name]
predictions = selected_model.predict(X_test)

# Use the actual classes instead of fixed Iris labels
labels = sorted(
    set(y_test.tolist()) | set(predictions.tolist()),
    key=str,
)

cm = confusion_matrix(
    y_test,
    predictions,
    labels=labels,
)

fig, ax = plt.subplots(figsize=(6, 4))

ConfusionMatrixDisplay(
    confusion_matrix=cm,
    display_labels=[str(label) for label in labels],
).plot(
    ax=ax,
    cmap="Blues",
    colorbar=False,
)

ax.set_title(selected_name)
st.pyplot(fig)
plt.close(fig)

    
# Feature Importance
if selected_name in ["Decision Tree", "Random Forest"]:
    st.subheader("🔍 Feature Importance")

    classifier = selected_model.named_steps["classifier"]
    preprocessor = selected_model.named_steps["preprocessor"]

    importances = classifier.feature_importances_
    feature_names = preprocessor.get_feature_names_out()

    importance_df = pd.DataFrame({
        "Feature": feature_names,
        "Importance": importances,
    }).sort_values("Importance", ascending=False)

    st.bar_chart(
        importance_df.head(15).set_index("Feature")
    )

# Live prediction
st.divider()
st.subheader("🔮 Try a New Prediction")

st.write(
    "Enter feature values and let the selected "
    "machine learning model make a prediction."
)

prediction_values = {}

col1, col2 = st.columns(2)
columns = [col1, col2]

for i, feature in enumerate(X.columns):
    with columns[i % 2]:
        if feature in numeric_features:
            valid_values = pd.to_numeric(
                X[feature], errors="coerce"
            ).dropna()

            if valid_values.empty:
                st.warning(
                    f"Feature '{feature}' contains no valid numeric values."
                )
                continue

            minimum = float(valid_values.min())
            maximum = float(valid_values.max())
            default = float(valid_values.median())

            # Allow editing even if all values are identical
            if minimum == maximum:
                minimum -= 1
                maximum += 1

            prediction_values[feature] = st.number_input(
                str(feature),
                min_value=minimum,
                max_value=maximum,
                value=min(max(default, minimum), maximum),
                key=f"prediction_{feature}",
            )

        else:
            options = X[feature].dropna().unique().tolist()

            if not options:
                st.warning(
                    f"Feature '{feature}' has no available values."
                )
                continue

            prediction_values[feature] = st.selectbox(
                str(feature),
                options=options,
                key=f"prediction_{feature}",
            )

prediction_model_name = st.selectbox(
    "Select a prediction model:",
    list(trained_models.keys()),
    key="prediction_model",
)

if st.button("🔮 Predict", key="run_prediction"):
    if len(prediction_values) != len(X.columns):
        st.error(
            "Some features have no usable values. "
            "Please check your dataset."
        )
    else:
        input_data = pd.DataFrame(
            [prediction_values],
            columns=X.columns,
        )

        model = trained_models[prediction_model_name]
        prediction = model.predict(input_data)[0]

        # Display original Iris species names when applicable
        if dataset_option == "Iris (Built-in)":
            predicted_label = iris.target_names[int(prediction)].title()
        else:
            predicted_label = str(prediction)

        st.success(f"Predicted Result: {predicted_label}")

        if hasattr(model, "predict_proba"):
            probabilities = model.predict_proba(input_data)[0]
            classes = model.named_steps["classifier"].classes_

            if dataset_option == "Iris (Built-in)":
                probability_labels = [
                    iris.target_names[int(label)].title()
                    for label in classes
                ]
            else:
                probability_labels = [str(label) for label in classes]

            probability_df = pd.DataFrame({
                "Class": probability_labels,
                "Probability": probabilities,
            }).set_index("Class")

            st.write("Prediction probabilities:")
            st.bar_chart(probability_df)